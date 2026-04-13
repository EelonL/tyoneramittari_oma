import io
import json
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font

APP_TITLE = "Työnerämittari"
DATA_DIR = Path(".runtime_data")
DATA_DIR.mkdir(exist_ok=True)


def now_local() -> datetime:
    return datetime.now().astimezone()


def fmt_ts(dt: datetime) -> str:
    return dt.isoformat(timespec="seconds")


def parse_ts(value: str) -> datetime:
    return datetime.fromisoformat(value)


def duration_seconds(start_iso: str, end_iso: str) -> float:
    return (parse_ts(end_iso) - parse_ts(start_iso)).total_seconds()


def human_duration(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def get_session_file() -> Path:
    if "session_file" not in st.session_state:
        session_id = f"session_{int(time.time() * 1000)}"
        st.session_state.session_file = str(DATA_DIR / f"{session_id}.json")
    return Path(st.session_state.session_file)


def default_state() -> Dict[str, Any]:
    return {
        "work_items": [],
        "events": [],
        "segments": [],
        "observations": [],
        "active_item": None,
        "active_start": None,
        "measurement_started": False,
        "measurement_finished": False,
        "measurement_label": "",
        "started_at": None,
        "finished_at": None,
    }


def recover_if_possible() -> None:
    files = sorted(DATA_DIR.glob("session_*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        return
    latest = files[0]
    try:
        recovered = json.loads(latest.read_text(encoding="utf-8"))
        if isinstance(recovered, dict) and recovered.get("events") is not None:
            st.session_state.app_state = {**default_state(), **recovered}
            st.session_state.session_file = str(latest)
    except Exception:
        pass


def ensure_state() -> None:
    if "app_state" not in st.session_state:
        st.session_state.app_state = default_state()
        recover_if_possible()


def persist_state() -> None:
    session_file = get_session_file()
    payload = dict(st.session_state.app_state)
    payload["last_saved_at"] = fmt_ts(now_local())
    session_file.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def add_event(event_type: str, item_name: Optional[str] = None, extra: Optional[Dict[str, Any]] = None) -> None:
    state = st.session_state.app_state
    payload = {"ts": fmt_ts(now_local()), "type": event_type, "item": item_name}
    if extra:
        payload.update(extra)
    state["events"].append(payload)
    persist_state()


def start_item(item_name: str) -> None:
    state = st.session_state.app_state
    ts = fmt_ts(now_local())
    if not state["measurement_started"]:
        state["measurement_started"] = True
        state["started_at"] = ts
    if state["active_item"] is not None and state["active_start"] is not None:
        previous = {
            "item": state["active_item"],
            "start": state["active_start"],
            "end": ts,
            "duration_seconds": duration_seconds(state["active_start"], ts),
        }
        state["segments"].append(previous)
        add_event("switch", item_name, {"previous_item": state["active_item"]})
    else:
        add_event("start", item_name)
    state["active_item"] = item_name
    state["active_start"] = ts
    state["measurement_finished"] = False
    persist_state()


def stop_measurement() -> None:
    state = st.session_state.app_state
    if state["active_item"] is None or state["active_start"] is None:
        return
    ts = fmt_ts(now_local())
    segment = {
        "item": state["active_item"],
        "start": state["active_start"],
        "end": ts,
        "duration_seconds": duration_seconds(state["active_start"], ts),
    }
    state["segments"].append(segment)
    state["finished_at"] = ts
    state["measurement_finished"] = True
    add_event("stop", state["active_item"])
    state["active_item"] = None
    state["active_start"] = None
    persist_state()


def reset_all() -> None:
    st.session_state.app_state = default_state()
    persist_state()


def segments_df() -> pd.DataFrame:
    state = st.session_state.app_state
    rows: List[Dict[str, Any]] = []
    for i, seg in enumerate(state["segments"], start=1):
        rows.append({
            "Järjestys": i,
            "Työnerä": seg["item"],
            "Alkuaika": seg["start"],
            "Loppuaika": seg["end"],
            "Kesto (s)": round(float(seg["duration_seconds"]), 1),
            "Kesto (hh:mm:ss)": human_duration(float(seg["duration_seconds"])),
        })
    if state["active_item"] and state["active_start"]:
        now_iso = fmt_ts(now_local())
        rows.append({
            "Järjestys": len(rows) + 1,
            "Työnerä": state["active_item"],
            "Alkuaika": state["active_start"],
            "Loppuaika": "KÄYNNISSÄ",
            "Kesto (s)": round(float(duration_seconds(state["active_start"], now_iso)), 1),
            "Kesto (hh:mm:ss)": human_duration(float(duration_seconds(state["active_start"], now_iso))),
        })
    return pd.DataFrame(rows)


def summary_df() -> pd.DataFrame:
    base = segments_df()
    if base.empty:
        return pd.DataFrame(columns=["Työnerä", "Toistot", "Yhteensä (s)", "Yhteensä (hh:mm:ss)"])
    finished_only = base[base["Loppuaika"] != "KÄYNNISSÄ"].copy()
    if finished_only.empty:
        return pd.DataFrame(columns=["Työnerä", "Toistot", "Yhteensä (s)", "Yhteensä (hh:mm:ss)"])
    grouped = (
        finished_only.groupby("Työnerä", as_index=False)
        .agg({"Kesto (s)": "sum", "Järjestys": "count"})
        .rename(columns={"Järjestys": "Toistot", "Kesto (s)": "Yhteensä (s)"})
    )
    grouped["Yhteensä (hh:mm:ss)"] = grouped["Yhteensä (s)"].apply(human_duration)
    grouped["Yhteensä (s)"] = grouped["Yhteensä (s)"].round(1)
    return grouped


def build_excel_bytes() -> bytes:
    detail = segments_df()
    summary = summary_df()
    meta = pd.DataFrame([
        {"Kenttä": "Mittauksen nimi", "Arvo": st.session_state.app_state.get("measurement_label", "")},
        {"Kenttä": "Mittaus aloitettu", "Arvo": st.session_state.app_state.get("started_at", "")},
        {"Kenttä": "Mittaus lopetettu", "Arvo": st.session_state.app_state.get("finished_at", "")},
        {"Kenttä": "Aktiivinen sessiotiedosto", "Arvo": str(get_session_file())},
    ])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        meta.to_excel(writer, index=False, sheet_name="Yhteenveto")
        start_row = len(meta) + 3
        summary.to_excel(writer, index=False, sheet_name="Yhteenveto", startrow=start_row)
        detail.to_excel(writer, index=False, sheet_name="Tapahtumat")

        obs_list = st.session_state.app_state.get("observations", [])
        if obs_list:
            obs_df = pd.DataFrame(obs_list).rename(columns={
                "ts": "Aikaleima",
                "item": "Työnerä",
                "joutuisuus": "Joutuisuus (%)",
                "huomio": "Huomio",
            })
            obs_df.to_excel(writer, index=False, sheet_name="Joutuisuus")
    output.seek(0)
    wb = load_workbook(output)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output.getvalue()


GLOBAL_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700;800&family=DM+Mono:wght@400;500&display=swap');

:root {
    --tts-blue: #1973ff;
    --tts-blue-dark: #0f5fe0;
    --tts-blue-soft: rgba(25, 115, 255, 0.07);
    --tts-border: rgba(25, 115, 255, 0.18);
    --tts-text: #0d2540;
    --tts-muted: #5a7490;
    --tts-stop: #c62828;
    --tts-stop-dark: #a31f1f;
    --tts-green: #00875a;
    --radius: 18px;
}

html, body, [class*="css"], .stApp, button, input, textarea {
    font-family: 'DM Sans', sans-serif !important;
}

/* ── Gradienttitausta ── */
[data-testid="stAppViewContainer"] {
    background: linear-gradient(145deg, #eef4ff 0%, #f8faff 55%, #f2eeff 100%);
    min-height: 100vh;
}
[data-testid="stHeader"],
[data-testid="stToolbar"] {
    background: transparent !important;
}
[data-testid="stSidebar"] {
    background: rgba(255,255,255,0.65);
    backdrop-filter: blur(14px);
    border-right: 1px solid var(--tts-border);
}

/* ── Ylätunniste ── */
.tts-app-title {
    font-size: 1.7rem;
    font-weight: 800;
    color: var(--tts-text);
    letter-spacing: -0.03em;
    line-height: 1.1;
}
.tts-app-subtitle {
    font-size: 0.92rem;
    color: var(--tts-muted);
    margin-top: 0.3rem;
    font-weight: 500;
}

/* ── Aktiivinen statusbanneri ── */
.active-banner {
    background: linear-gradient(90deg, #1973ff 0%, #4f9dff 100%);
    color: white;
    border-radius: var(--radius);
    padding: 0.9rem 1.3rem;
    font-weight: 700;
    font-size: 1.05rem;
    margin-bottom: 1rem;
    box-shadow: 0 6px 24px rgba(25,115,255,0.28);
    display: flex;
    align-items: center;
    gap: 0.6rem;
    letter-spacing: -0.01em;
}
.active-banner .timer {
    font-family: 'DM Mono', monospace;
    font-size: 1.1rem;
    background: rgba(255,255,255,0.2);
    border-radius: 8px;
    padding: 0.1rem 0.55rem;
    margin-left: auto;
}

/* ── Työnerapainikkeet (glassmorphism) ── */
.work-items-section div[data-testid="stButton"] > button {
    width: 100%;
    min-height: 74px;
    border-radius: var(--radius);
    text-align: left;
    font-size: 1.05rem;
    font-weight: 600;
    padding: 1rem 1.2rem;
    margin-bottom: 0.5rem;
    border: 1px solid var(--tts-border) !important;
    background: rgba(255, 255, 255, 0.68) !important;
    backdrop-filter: blur(12px);
    color: var(--tts-text) !important;
    box-shadow: 0 2px 12px rgba(25,115,255,0.07);
    transition: transform 0.13s ease, box-shadow 0.13s ease, background 0.13s ease;
    letter-spacing: -0.01em;
}
.work-items-section div[data-testid="stButton"] > button:hover {
    transform: translateY(-2px);
    box-shadow: 0 8px 24px rgba(25,115,255,0.16);
    background: rgba(255, 255, 255, 0.9) !important;
    border-color: var(--tts-blue) !important;
}

/* ── Aktiivinen työnerapainike ── */
@keyframes pulse-border {
    0%, 100% { box-shadow: 0 0 0 0 rgba(25,115,255,0.45), 0 6px 20px rgba(25,115,255,0.25); }
    50%       { box-shadow: 0 0 0 7px rgba(25,115,255,0), 0 6px 20px rgba(25,115,255,0.25); }
}
.active-work-item div[data-testid="stButton"] > button {
    background: linear-gradient(135deg, #1973ff 0%, #2e87ff 100%) !important;
    color: white !important;
    border: 1px solid rgba(255,255,255,0.25) !important;
    min-height: 90px;
    animation: pulse-border 2.2s ease infinite;
    font-size: 1.1rem;
}
.active-work-item div[data-testid="stButton"] > button:hover {
    transform: translateY(-2px);
    background: linear-gradient(135deg, #0f5fe0 0%, #1973ff 100%) !important;
}

/* ── Lopetuspainike ── */
.stop-section div[data-testid="stButton"] > button {
    min-height: 68px;
    border-radius: var(--radius);
    background: linear-gradient(135deg, var(--tts-stop) 0%, #e53935 100%) !important;
    color: white !important;
    border: 1px solid rgba(255,255,255,0.15) !important;
    text-align: center;
    font-weight: 700;
    font-size: 1.05rem;
    box-shadow: 0 4px 16px rgba(198,40,40,0.28);
    transition: transform 0.13s ease, box-shadow 0.13s ease;
}
.stop-section div[data-testid="stButton"] > button:hover {
    transform: translateY(-2px);
    box-shadow: 0 8px 24px rgba(198,40,40,0.35);
}

/* ── Metriikat ── */
[data-testid="metric-container"] {
    background: rgba(255,255,255,0.7);
    backdrop-filter: blur(10px);
    border: 1px solid var(--tts-border);
    border-radius: var(--radius);
    padding: 1rem 1.2rem;
    box-shadow: 0 2px 12px rgba(25,115,255,0.07);
}
[data-testid="metric-container"] [data-testid="stMetricValue"] {
    font-family: 'DM Mono', monospace !important;
    font-size: 1.5rem !important;
    color: var(--tts-blue) !important;
    font-weight: 600 !important;
}
[data-testid="metric-container"] [data-testid="stMetricLabel"] {
    font-weight: 600;
    font-size: 0.82rem;
    text-transform: uppercase;
    letter-spacing: 0.05em;
    color: var(--tts-muted);
}

/* ── Osioiden otsikot ── */
h2, h3 {
    color: var(--tts-text) !important;
    letter-spacing: -0.02em !important;
}

/* ── Divider ── */
hr {
    border-color: var(--tts-border) !important;
    margin: 1.5rem 0 !important;
}

/* ── Mobiili ── */
@media (max-width: 768px) {
    .block-container { padding-left: 0.5rem !important; padding-right: 0.5rem !important; }
    .tts-app-title { font-size: 1.35rem; }
}
</style>
"""


def render_brand_header() -> None:
    logo_path = next((p for p in ["Logo.jpg", "Logo.png"] if Path(p).exists()), None)

    st.markdown(GLOBAL_CSS, unsafe_allow_html=True)

    # Sivupalkki: toimintavarmuustiedot
    with st.sidebar:
        st.markdown("### ⚙️ Tietoa sovelluksesta")
        st.markdown(
            """
            - Jokainen aloitus, vaihto ja lopetus tallennetaan heti JSON-palautustiedostoon.
            - Lyhyt verkkokatkos ei yleensä riko mittausta, jos selainvälilehti pysyy auki.
            - Täydellinen selainpäivitys tai sessioiden katkeaminen Streamlit Cloudissa voi silti katkaista tilan.
            - Tuotantoversiossa suositellaan lisäksi selaimen localStorage-varmistusta tai taustatietokantaa.
            """
        )

    c1, c2 = st.columns([5, 1])
    with c1:
        st.markdown(
            """
            <div class="tts-app-title">⏱️ Työnerämittari</div>
            <div class="tts-app-subtitle">Käynnistä, vaihda ja lopeta työnerät yhdellä näkymällä</div>
            """,
            unsafe_allow_html=True,
        )
    with c2:
        if logo_path:
            st.image(logo_path, use_container_width=True)

    st.write("Mittaa työnerien alkamis- ja päättymisajat sekä lataa tulokset Excel-tiedostona.")


def work_items_editor() -> None:
    st.subheader("1. Määritä työnerät")
    state = st.session_state.app_state

    state["measurement_label"] = st.text_input(
        "Mittauksen nimi",
        value=state.get("measurement_label", ""),
        placeholder="Esim. Huoneiston A123 työvaiheiden seuranta",
    )

    initial = "\n".join(state.get("work_items", []))
    text = st.text_area(
        "Syötä yksi työnerä per rivi",
        value=initial,
        height=160,
        placeholder="Esim.\nMuottityö\nRaudoitus\nBetonointi\nSiivous",
    )

    items = [row.strip() for row in text.splitlines() if row.strip()]
    state["work_items"] = items
    persist_state()

    c1, c2 = st.columns(2)
    with c1:
        if st.button("Tallenna työnerät", use_container_width=True):
            persist_state()
            st.success("Työnerät tallennettu.")
    with c2:
        if st.button("Tyhjennä kaikki", use_container_width=True):
            reset_all()
            st.rerun()


@st.fragment(run_every="1s")
def measurement_ui() -> None:
    st.subheader("2. Käynnistä mittaus")
    state = st.session_state.app_state

    if not state["work_items"]:
        st.warning("Lisää ensin vähintään yksi työnerä.")
        return

    active_item = state.get("active_item")
    active_start = state.get("active_start")
    active_elapsed = ""

    if active_item and active_start:
        elapsed_seconds = (now_local() - parse_ts(active_start)).total_seconds()
        active_elapsed = human_duration(elapsed_seconds)

        # ── Aktiivinen statusbanneri ──
        st.markdown(
            f"""
            <div class="active-banner">
                🟢 Käynnissä: <strong>{active_item}</strong>
                <span class="timer">{active_elapsed}</span>
            </div>
            """,
            unsafe_allow_html=True,
        )

    # ── Metriikat ──
    summary = summary_df()
    total_segments = len(state["segments"])
    total_seconds = summary["Yhteensä (s)"].sum() if not summary.empty else 0.0
    unique_items = summary["Työnerä"].nunique() if not summary.empty else 0

    m1, m2, m3 = st.columns(3)
    m1.metric("Mittausaika yhteensä", human_duration(float(total_seconds)))
    m2.metric("Vaihtoja yhteensä", str(total_segments))
    m3.metric("Eri työnerät", str(unique_items))

    st.caption("Valitse aktiivinen työnerä alla. Lisää työneriä löytyy rullaamalla alaspäin.")
    st.markdown('<div class="work-items-section">', unsafe_allow_html=True)

    for idx, item in enumerate(state["work_items"], start=1):
        is_active = active_item == item
        label = f"{idx}. {item}"

        if is_active:
            label = f"🟢 {idx}. {item}  ·  {active_elapsed}"
            st.markdown('<div class="active-work-item">', unsafe_allow_html=True)
            if st.button(label, key=f"item_{idx}", use_container_width=True):
                start_item(item)
                st.rerun()
            st.markdown("</div>", unsafe_allow_html=True)
        else:
            if st.button(label, key=f"item_{idx}", use_container_width=True):
                start_item(item)
                st.rerun()

    st.markdown("</div>", unsafe_allow_html=True)

    # ── Pikalisäys: uusi työnerä lennosta ──
    st.markdown(
        """
        <div style="margin-top:1.2rem; padding: 1rem 1.2rem;
            background: rgba(255,255,255,0.55); backdrop-filter: blur(10px);
            border: 1px dashed rgba(25,115,255,0.35); border-radius: 18px;">
            <div style="font-size:0.82rem; font-weight:600; color:#5a7490;
                text-transform:uppercase; letter-spacing:0.06em; margin-bottom:0.6rem;">
                ＋ Lisää puuttuva työnerä
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    col_input, col_btn = st.columns([4, 1])
    with col_input:
        new_item = st.text_input(
            "Uusi työnerä",
            key="quick_add_input",
            placeholder="Kirjoita työnerän nimi...",
            label_visibility="collapsed",
        )
    with col_btn:
        if st.button("Lisää ja käynnistä", key="quick_add_btn", use_container_width=True):
            name = new_item.strip()
            if name and name not in state["work_items"]:
                state["work_items"].append(name)
                persist_state()
                start_item(name)
                st.rerun()
            elif name in state["work_items"]:
                start_item(name)
                st.rerun()
            else:
                st.warning("Anna työnerän nimi.")


@st.fragment()
def observation_ui() -> None:
    state = st.session_state.app_state
    active_item = state.get("active_item")

    if not active_item:
        return

    st.markdown(
        """
        <div style="margin: 0.6rem 0 0.5rem;
            padding: 1rem 1.2rem 0.7rem;
            background: rgba(255,255,255,0.55); backdrop-filter: blur(10px);
            border: 1px solid rgba(25,115,255,0.18); border-radius: 18px;">
            <div style="font-size:0.82rem; font-weight:600; color:#5a7490;
                text-transform:uppercase; letter-spacing:0.06em; margin-bottom:0.7rem;">
                📋 Joutuisuushavainto — käynnissä olevalle työnerälle
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.caption(f"Kirjataan kohteelle: **{active_item}**")

    col_custom, col_note = st.columns([1, 2])
    with col_custom:
        custom_pace = st.number_input(
            "Joutuisuus %",
            min_value=10, max_value=200,
            value=st.session_state.get("obs_pace_value", 100),
            step=5,
            key="obs_pace_custom",
        )
        st.session_state.obs_pace_value = custom_pace

    note_key_idx = st.session_state.get("obs_note_idx", 0)
    with col_note:
        note = st.text_input(
            "Vapaaehtoinen huomio",
            key=f"obs_note_{note_key_idx}",
            placeholder="Esim. häiriötekijä, erikoistilanne...",
        )

    if st.button("📋 Kirjaa havainto", key="obs_save_btn", use_container_width=True):
        obs = {
            "ts": fmt_ts(now_local()),
            "item": active_item,
            "joutuisuus": st.session_state.get("obs_pace_value", 100),
            "huomio": note.strip(),
        }
        state["observations"].append(obs)
        persist_state()
        st.session_state.obs_note_idx = note_key_idx + 1
        st.success(
            f"✓ Kirjattu — {active_item}, joutuisuus "
            f"{obs['joutuisuus']}%"
            + (f", huomio: {obs['huomio']}" if obs['huomio'] else "")
        )


def stop_button_ui() -> None:
    state = st.session_state.app_state
    active_item = state.get("active_item")
    active_start = state.get("active_start")

    st.markdown("---")
    st.subheader("3. Mittauksen lopetus")

    stop_label = "🛑 Lopeta mittaus"
    if active_item:
        stop_label = f"🛑 Lopeta: {active_item}"

    st.markdown('<div class="stop-section">', unsafe_allow_html=True)
    if st.button(stop_label, key="stop_measurement_main", use_container_width=True):
        if active_item and active_start:
            stop_measurement()
            st.rerun()
        else:
            st.warning("Aktiivista työnerää ei ole käynnissä.")
    st.markdown("</div>", unsafe_allow_html=True)


def live_tables() -> None:
    detail = segments_df()
    summary = summary_df()

    st.subheader("Mittausdata")
    t1, t2, t3 = st.tabs(["Tapahtumat", "Yhteenveto", "Joutuisuus"])

    with t1:
        if detail.empty:
            st.write("Ei vielä rivejä.")
        else:
            st.dataframe(detail, use_container_width=True)

    with t2:
        if summary.empty:
            st.write("Ei vielä yhteenvetoa.")
        else:
            st.dataframe(summary, use_container_width=True)

    with t3:
        obs_list = st.session_state.app_state.get("observations", [])
        if not obs_list:
            st.write("Ei vielä havaintoja.")
        else:
            obs_df = pd.DataFrame(obs_list).rename(columns={
                "ts": "Aikaleima",
                "item": "Työnerä",
                "joutuisuus": "Joutuisuus (%)",
                "huomio": "Huomio",
            })
            st.dataframe(obs_df, use_container_width=True)


def save_ui() -> None:
    st.subheader("4. Excel-tiedoston muodostus")
    filename = f"tyoneramittaus_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    excel_bytes = build_excel_bytes()

    st.download_button(
        "⬇️ Lataa Excel",
        data=excel_bytes,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    st.caption("Excel ladataan käyttäjän laitteelle.")


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="⏱️", layout="wide")
    ensure_state()

    render_brand_header()
    work_items_editor()
    st.divider()
    measurement_ui()
    observation_ui()
    stop_button_ui()
    st.divider()
    live_tables()
    st.divider()
    save_ui()


if __name__ == "__main__":
    main()
